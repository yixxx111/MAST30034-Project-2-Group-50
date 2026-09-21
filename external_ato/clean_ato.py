"""Clean ATO 2021-22 Table 6B without changing the official workbook.

Run from the repository root: python external_ato/clean_ato.py --help
Outputs describe areas retrospectively; they are not individual credit features.
"""
from __future__ import annotations

import argparse
import collections
import csv
import hashlib
import io
import json
import math
import re
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

SOURCE_URL = "https://data.gov.au/data/dataset/5fa69f19-ec44-4c46-88eb-0f6fd5c2f43b/resource/00847743-8b5b-4d41-a15a-1a28cf00ea81/download/ts22individual06taxablestatusstatesa4postcode.xlsx"
CATALOG_URL = "https://data.gov.au/data/dataset/taxation-statistics-postcode-data"
RAW_NAME = "ts22individual06taxablestatusstatesa4postcode.xlsx"
STATES = {"ACT", "NSW", "NT", "QLD", "SA", "TAS", "VIC", "WA"}
# State-other aggregates are not individual postcodes; Notes do not establish every exclusion cause.
STATE_OTHER = re.compile(r"(ACT|NSW|NT|QLD|SA|TAS|VIC|WA) other")
# Project review threshold for the denominators of derived means/shares; not an ATO rule, never deletes rows.
SMALL_DENOMINATOR = 100
# Australia Post ranges reserved for PO boxes / large volume receivers. Heuristic label for audit only.
PO_BOX_LVR_RANGES = ((200, 299), (900, 999), (1000, 1999), (5800, 5999), (6800, 6999),
                     (7800, 7999), (8000, 8999), (9000, 9999))
# Exact headers are intentional: a different source edition must not be silently accepted.
FIELDS = {
    "ato_individual_count": ("Individuals\nno.", "count", "All individuals in this Table 6B area record"),
    "ato_taxable_income_reporter_count": ("Taxable income or loss4\nno.", "count", "Count for the taxable income or loss label; not the total individual count"),
    "ato_taxable_income_or_loss_total": ("Taxable income or loss4\n$", "AUD/year", "Published total taxable income or loss"),
    "ato_salary_recipient_count": ("Salary or wages\nno.", "count", "Count for salary or wages label"),
    "ato_salary_or_wages_total": ("Salary or wages\n$", "AUD/year", "Published total salary or wages"),
    "ato_net_tax_count": ("Net tax\nno.", "count", "Count for net tax label; not a fraud or default count"),
    "ato_net_tax_total": ("Net tax\n$", "AUD/year", "Published net tax amount"),
}
DERIVED = {
    "ato_taxable_income_or_loss_per_reporter": ("ato_taxable_income_or_loss_total", "ato_taxable_income_reporter_count", "AUD/year", "Derived total divided by corresponding label count, not median or disposable income"),
    "ato_salary_or_wages_per_recipient": ("ato_salary_or_wages_total", "ato_salary_recipient_count", "AUD/year", "Derived salary total divided by salary recipient count"),
    "ato_salary_recipient_share": ("ato_salary_recipient_count", "ato_individual_count", "proportion 0-1", "Salary recipients divided by all individuals; not a population employment rate"),
    "ato_net_tax_payer_share": ("ato_net_tax_count", "ato_individual_count", "proportion 0-1", "Individuals with a net tax amount divided by all individuals; not a compliance or risk measure"),
}


def postcode(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value != int(value):
            return None
        value = str(int(value))
    value = str(value).strip()
    if not re.fullmatch(r"[0-9]{1,4}", value):
        return None
    value = value.zfill(4)
    # Australian postcodes start at 0200; "5" -> "0005" style values are not postcodes.
    return None if int(value) < 200 else value


def po_box_lvr_range(code):
    return any(low <= int(code) <= high for low, high in PO_BOX_LVR_RANGES)


def number(value, count=False):
    """Blanks become missing. Unexpected symbols fail until source notes are reviewed."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "blank"
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Unrecognised numeric value {value!r}; review source notes before mapping")
    if not math.isfinite(value):
        return None, "nonfinite"
    if count and (value < 0 or value != int(value)):
        return None, "invalid_count"
    return (int(value) if count else value), None


def divide(numerator, denominator, share=False):
    if numerator is None or denominator is None:
        return None, "missing_input"
    if denominator <= 0:
        return None, "nonpositive_denominator"
    result = numerator / denominator
    if share and not 0 <= result <= 1:
        return None, "share_out_of_range"
    return result, None


def clean(headers, source_rows):
    if len(headers) != len(set(headers)):
        raise ValueError("Duplicate source column headers")
    required = {"State/ Territory1", "Statistical Area Level 4 (SA4)2", "Postcode"}
    required.update(v[0] for v in FIELDS.values())
    if missing := required - set(headers):
        raise ValueError(f"Missing required headers: {sorted(missing)}")
    result, excluded, issues, seen = [], [], [], set()
    for excel_row, values in enumerate(source_rows, start=3):
        if not any(v is not None for v in values):
            continue
        row = dict(zip(headers, values))
        key = postcode(row["Postcode"])
        state = "" if row["State/ Territory1"] is None else str(row["State/ Territory1"]).strip()
        if key is None:
            label = "" if row["Postcode"] is None else str(row["Postcode"]).strip()
            if STATE_OTHER.fullmatch(label) and label == f"{state} other" and state in STATES:
                reason = "state_other_aggregate"
            elif label == "Overseas" and state == "Overseas":
                reason = "overseas_aggregate"
            else:
                raise ValueError(f"Excel row {excel_row}: unrecognised postcode label {row['Postcode']!r}; review before excluding")
            excluded.append({"excel_row": excel_row, "source_postcode": label, "source_state": state, "reason": reason,
                             "individual_count": number(row[FIELDS["ato_individual_count"][0]], count=True)[0],
                             "taxable_income_or_loss_total": number(row[FIELDS["ato_taxable_income_or_loss_total"][0]])[0]})
            continue
        if state not in STATES:
            raise ValueError(f"Excel row {excel_row}: postcode {key} has unexpected state {state!r}; not silently dropped")
        if key in seen:
            raise ValueError(f"Duplicate normalised postcode: {key}; do not arbitrarily deduplicate")
        seen.add(key)
        sa4 = "" if row["Statistical Area Level 4 (SA4)2"] is None else str(row["Statistical Area Level 4 (SA4)2"]).strip()
        item = {"postcode": key, "ato_income_year": "2021-22", "ato_state": state, "ato_sa4_name": sa4,
                "ato_sa4_is_state_other": bool(STATE_OTHER.fullmatch(sa4))}
        issue_count = len(issues)
        for field, (source, unit, _) in FIELDS.items():
            item[field], reason = number(row[source], count=unit == "count")
            if reason:
                issues.append({"postcode": key, "excel_row": excel_row, "field": field,
                               "source_value": row[source], "reason": reason})
        for field, (num, den, unit, _) in DERIVED.items():
            item[field], reason = divide(item[num], item[den], share=unit == "proportion 0-1")
            if reason:
                issues.append({"postcode": key, "excel_row": excel_row, "field": field,
                               "source_value": None, "reason": reason})
        for field, (_, unit, _) in FIELDS.items():
            if unit == "count" and field != "ato_individual_count" and item[field] is not None and item['ato_individual_count'] is not None:
                if item[field] > item['ato_individual_count']:
                    issues.append({"postcode": key, "excel_row": excel_row, "field": field,
                                   "source_value": item[field], "reason": "label_count_above_individual_count_retained"})
        denominators = [item[den] for _, den, _, _ in DERIVED.values() if item[den] is not None]
        item["ato_min_derived_denominator"] = min(denominators) if denominators else None
        item["ato_small_denominator"] = min(denominators) < SMALL_DENOMINATOR if denominators else None
        item["ato_nonpositive_taxable_income_total"] = item["ato_taxable_income_or_loss_total"] <= 0 if item["ato_taxable_income_or_loss_total"] is not None else None
        item["ato_any_feature_missing"] = any(item[k] is None for k in [*FIELDS, *DERIVED])
        item["ato_quality_issue"] = len(issues) > issue_count
        result.append(item)
    if not result:
        raise ValueError("No valid domestic postcode rows")
    return sorted(result, key=lambda r: r['postcode']), excluded, issues


def consumer_counts(path):
    """Read postcode and state only; never export names, addresses or IDs.

    Returns a Counter keyed by (normalised postcode or None, consumer state).
    """
    def read(stream):
        rows = csv.DictReader(stream, delimiter="|")
        if not {"postcode", "state"} <= set(rows.fieldnames or []):
            raise ValueError("Expected pipe-delimited tbl_consumer.csv with postcode and state headers")
        return collections.Counter((postcode(r['postcode']), (r['state'] or '').strip()) for r in rows)
    if path.suffix.lower() == '.zip':
        with zipfile.ZipFile(path) as z:
            matches = [n for n in z.namelist() if n.endswith('/tbl_consumer.csv') or n == 'tbl_consumer.csv']
            if len(matches) != 1:
                raise ValueError("Expected exactly one tbl_consumer.csv in ZIP")
            with z.open(matches[0]) as f:
                return read(io.TextIOWrapper(f, encoding='utf-8-sig'))
    with path.open(encoding='utf-8-sig', newline='') as f:
        return read(f)


def other_source_postcodes(spec):
    """Parse NAME=cleaned.csv (needs a postcode column) for the joint coverage audit."""
    name, _, file = spec.partition('=')
    if not name or not file:
        raise ValueError(f"--compare expects NAME=path.csv, got {spec!r}")
    with Path(file).open(encoding='utf-8-sig', newline='') as f:
        rows = csv.DictReader(f)
        if 'postcode' not in (rows.fieldnames or []):
            raise ValueError(f"{file} has no postcode column")
        return name, {postcode(r['postcode']) for r in rows} - {None}


def coverage(pair_counts, rows, others=None):
    """Consumer-row and distinct-postcode coverage, in the same layout as the Census/SEIFA modules."""
    lookup = {r['postcode']: r for r in rows}
    counts = collections.Counter()
    for (key, _), n in pair_counts.items():
        counts[key] += n
    total = sum(counts.values())
    status_of = lambda key: 'missing_or_invalid_postcode' if key is None else ('matched' if key in lookup else 'postcode_not_in_ato')
    totals, distinct, exceptions = collections.Counter(), collections.Counter(), []
    for key, count in counts.items():
        status = status_of(key)
        totals[status] += count
        if key is not None:
            distinct[status] += 1
        if status != 'matched':
            exceptions.append({'postcode': key, 'consumer_count': count, 'status': status,
                               'postcode_range_class': None if key is None else
                               ('heuristic_special_range' if po_box_lvr_range(key) else 'outside_heuristic_special_range')})
    exceptions.sort(key=lambda r: (r['postcode'] is None, r['postcode'] or ''))

    def line(scope, metric, n, denominator):
        return {'scope': scope, 'metric': metric, 'rows': n, 'denominator': denominator,
                'rate': n / denominator if denominator else None}
    report = [line('consumer_rows', k, totals[k], total)
              for k in ['matched', 'postcode_not_in_ato', 'missing_or_invalid_postcode']]
    valid = sum(distinct.values())
    report += [line('distinct_valid_postcodes', k, distinct[k], valid) for k in ['matched', 'postcode_not_in_ato']]
    mismatch = collections.Counter()
    for (key, state), n in pair_counts.items():
        if key in lookup and state != lookup[key]['ato_state']:
            mismatch[(state, lookup[key]['ato_state'])] += n
    report.append(line('consumer_rows', 'matched_but_consumer_state_differs_from_ato_state', sum(mismatch.values()), total))
    state_report = [{'consumer_state': a, 'ato_state': b, 'consumer_rows': n} for (a, b), n in sorted(mismatch.items())]
    if others:
        every = set(lookup)
        for name, codes in others.items():
            every &= codes
            report.append(line('consumer_rows', f'postcode_in_ato_and_{name}',
                               sum(n for k, n in counts.items() if k in lookup and k in codes), total))
        report.append(line('consumer_rows', 'postcode_in_all_sources', sum(n for k, n in counts.items() if k in every), total))
    missingness = []
    for field in [*FIELDS, *DERIVED]:
        available = sum(n for key, n in counts.items() if key in lookup and lookup[key][field] is not None)
        missingness.append({'field': field, 'available_consumer_rows': available,
                            'missing_consumer_rows': total - available, 'total_consumer_rows': total})
    return report, exceptions, missingness, state_report


def reconcile_table_6a(table_a, table_b):
    """Table 6A (taxable + non-taxable) must add up to Table 6B for every selected source column."""
    head_a, head_b = list(table_a[1]), list(table_b[1])
    pa, pb = head_a.index('Postcode'), head_b.index('Postcode')
    report = []
    for field, (source, _, _) in FIELDS.items():
        ja, jb = head_a.index(source), head_b.index(source)
        sums = collections.Counter()
        for r in table_a[2:]:
            if r[pa] is not None:
                sums[str(r[pa])] += r[ja] or 0
        published = {str(r[pb]): r[jb] or 0 for r in table_b[2:] if r[pb] is not None}
        common = set(sums) & set(published)
        differences = [abs(sums[k] - published[k]) for k in common]
        # $ totals are rounded per component, so allow 1 per taxable-status component.
        tolerance = 0 if source.endswith('no.') else 2
        report.append({'field': field, 'records_compared': len(common),
                       'records_only_in_6a': len(set(sums) - set(published)),
                       'records_only_in_6b': len(set(published) - set(sums)),
                       'records_beyond_tolerance': sum(d > tolerance for d in differences),
                       'max_abs_difference': max(differences, default=0), 'tolerance': tolerance})
    return report


def assert_reconciliation(report):
    if any(r['records_beyond_tolerance'] or r['records_only_in_6a'] or r['records_only_in_6b'] for r in report):
        raise ValueError('Table 6A/6B reconciliation failed; no clean output written')


def write_parquet(path, rows):
    """Explicit nullable schema preserves postcodes, integer counts and boolean flags."""
    strings = {'postcode', 'ato_income_year', 'ato_state', 'ato_sa4_name'}
    flags = {'ato_sa4_is_state_other', 'ato_small_denominator', 'ato_nonpositive_taxable_income_total',
             'ato_any_feature_missing', 'ato_quality_issue'}
    integers = {k for k, (_, unit, _) in FIELDS.items() if unit == 'count'} | {'ato_min_derived_denominator'}
    schema = pa.schema([(k, pa.string() if k in strings else pa.bool_() if k in flags else
                         pa.int64() if k in integers else pa.float64()) for k in rows[0]])
    table = pa.Table.from_pylist(rows, schema=schema)
    pq.write_table(table, path, compression='snappy')
    if not pq.read_table(path).equals(table):
        raise ValueError('Parquet round-trip changed values or schema')


def write_csv(path, rows, fields=None):
    fields = fields or list(rows[0])
    with path.open('w', encoding='utf-8', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def run(args):
    repo = Path(__file__).resolve().parent.parent
    raw = (args.xlsx or repo / 'tables/external/ato_2021_22' / RAW_NAME).resolve()
    out = (args.output or Path(__file__).resolve().parent / 'results').resolve()
    if out == raw.parent or out in raw.parents:
        raise ValueError('Output cannot contain the raw input')
    if out.exists() and any(out.iterdir()):
        raise ValueError('Output must be empty; choose a new --output folder for a rerun')
    book = openpyxl.load_workbook(raw, read_only=True, data_only=True)
    if 'Table 6B' not in book.sheetnames or 'Notes' not in book.sheetnames:
        raise ValueError('Expected ATO Table 6B and Notes sheets')
    source = list(book['Table 6B'].values)
    if '2021–22' not in str(source[0][0]):
        raise ValueError('Unexpected income year in workbook title')
    rows, excluded, issues = clean(source[1], source[2:])
    reconciliation = reconcile_table_6a(list(book['Table 6A'].values), source) if 'Table 6A' in book.sheetnames else None
    if reconciliation is not None:
        assert_reconciliation(reconciliation)
    notes = '\n'.join(' | '.join(str(v) for v in r if v is not None) for r in book['Notes'].values)
    book.close()
    others = dict(other_source_postcodes(spec) for spec in args.compare or [])
    if others and not args.consumers:
        raise ValueError('--compare needs --consumers')
    reports = coverage(consumer_counts(args.consumers), rows, others) if args.consumers else None
    out.mkdir(parents=True, exist_ok=True)
    write_csv(out / 'ato_clean.csv', rows)
    write_parquet(out / 'ato_clean.parquet', rows)
    write_csv(out / 'excluded_records.csv', excluded, ['excel_row', 'source_postcode', 'source_state', 'reason',
                                                       'individual_count', 'taxable_income_or_loss_total'])
    if reconciliation:
        write_csv(out / 'table6a_reconciliation.csv', reconciliation)
    write_csv(out / 'quality_issues.csv', issues, ['postcode', 'excel_row', 'field', 'source_value', 'reason'])
    dictionary = [{'field': k, 'source_or_formula': h.replace('\n', ' '), 'unit': u, 'definition': d}
                  for k, (h, u, d) in FIELDS.items()]
    dictionary += [{'field': k, 'source_or_formula': f'{n} / {d}', 'unit': u, 'definition': desc}
                   for k, (n, d, u, desc) in DERIVED.items()]
    descriptions = {
        'postcode': ('Postcode', 'string', 'Four digits; format-valid does not establish postal validity'),
        'ato_income_year': ('Workbook title', 'year range', '2021-22 income year; retrospective context'),
        'ato_state': ('State/ Territory1', 'category', 'Source state; postcode-based ATO allocation'),
        'ato_sa4_name': ('Statistical Area Level 4 (SA4)2', 'category', 'Source SA4 name; largest-weight correspondence, not precise consumer location; placeholder when ato_sa4_is_state_other'),
        'ato_sa4_is_state_other': ('ato_sa4_name matches "<STATE> other"', 'boolean', 'Source uses a state-other placeholder; cause is not inferred from the label'),
        'ato_min_derived_denominator': ('min of the denominators used by derived fields', 'count', 'Smallest count behind any derived mean/share for this postcode'),
        'ato_small_denominator': (f'ato_min_derived_denominator < {SMALL_DENOMINATOR}', 'boolean', 'Project review threshold only; derived means on few individuals are unstable; no row deletion'),
        'ato_nonpositive_taxable_income_total': ('ato_taxable_income_or_loss_total <= 0', 'boolean', 'Review flag; source value retained'),
        'ato_any_feature_missing': ('Selected numeric fields', 'boolean', 'Any selected or derived numeric feature unavailable'),
        'ato_quality_issue': ('quality_issues.csv', 'boolean', 'At least one recorded numeric/ratio/count issue'),
    }
    dictionary += [{'field': k, 'source_or_formula': s, 'unit': u, 'definition': d} for k, (s, u, d) in descriptions.items()]
    write_csv(out / 'data_dictionary.csv', dictionary)
    profile = []
    for field in [*FIELDS, *DERIVED]:
        values = sorted(r[field] for r in rows if r[field] is not None)
        profile.append({'field': field, 'rows': len(rows), 'missing': len(rows)-len(values),
                        'minimum': values[0] if values else None, 'maximum': values[-1] if values else None,
                        'zero_count': sum(v == 0 for v in values), 'negative_count': sum(v < 0 for v in values)})
    write_csv(out / 'feature_quality_profile.csv', profile)
    if reports:
        write_csv(out / 'consumer_join_coverage.csv', reports[0])
        write_csv(out / 'consumer_postcode_exceptions.csv', reports[1], ['postcode', 'consumer_count', 'status', 'postcode_range_class'])
        write_csv(out / 'consumer_feature_missingness.csv', reports[2])
        write_csv(out / 'consumer_state_consistency.csv', reports[3], ['consumer_state', 'ato_state', 'consumer_rows'])
    (out / 'source_notes.txt').write_text(notes, encoding='utf-8')
    metadata = {
        'source_url': SOURCE_URL, 'catalog_url': CATALOG_URL, 'license': 'CC BY 2.5 Australia',
        'source_file': raw.name, 'source_sha256': hashlib.sha256(raw.read_bytes()).hexdigest(),
        'worksheet': 'Table 6B', 'income_year': '2021-22',
        'source_processing_cutoff': '2023-10-31',
        'catalog_resource_created': '2024-08-15T00:17:06.116148 (not asserted to be first publication date)',
        'intended_use': 'Retrospective area context only; not available during 2021-22 transactions',
        'run_utc': datetime.now(timezone.utc).isoformat(), 'python': sys.version.split()[0], 'openpyxl': openpyxl.__version__,
        'pyarrow': pa.__version__, 'parquet_postcode_type': 'string',
        'comparison_sources': {name: {'path': str(Path(spec.partition('=')[2])),
                                     'sha256': hashlib.sha256(Path(spec.partition('=')[2]).read_bytes()).hexdigest()}
                               for spec in args.compare or [] for name in [spec.partition('=')[0]]},
        'source_data_rows': sum(any(v is not None for v in r) for r in source[2:]),
        'clean_rows': len(rows), 'clean_columns': len(rows[0]), 'excluded_rows': len(excluded), 'quality_issue_cells': len(issues),
        'source_individuals_total': sum(r['ato_individual_count'] or 0 for r in rows) + sum(e['individual_count'] or 0 for e in excluded),
        'excluded_individuals_by_reason': {reason: sum(e['individual_count'] or 0 for e in excluded if e['reason'] == reason)
                                           for reason in sorted({e['reason'] for e in excluded})},
        'observed_min_individual_count': min(r['ato_individual_count'] for r in rows if r['ato_individual_count'] is not None),
        'small_denominator_threshold': SMALL_DENOMINATOR,
        'small_denominator_postcodes': sum(bool(r['ato_small_denominator']) for r in rows),
        'sa4_state_other_postcodes': sum(r['ato_sa4_is_state_other'] for r in rows),
        'table6a_reconciliation_records_beyond_tolerance': sum(r['records_beyond_tolerance'] for r in reconciliation) if reconciliation else 'not_run',
        'consumer_coverage': reports[0] if reports else 'not_run_no_consumer_input',
        'transaction_enrichment': 'not_run',
        'policies': ['No imputation, trimming, winsorisation or statistical outlier deletion',
                     'Negative income/loss amounts retained; invalid counts become null and are audited',
                     'Unrecognised text numeric markers fail pending source review',
                     'Means use each corresponding label count; no median inferred',
                     'Total income label not selected; avoid treating correlated income measures as independent evidence',
                     'State-other and Overseas aggregates excluded with separate reasons; unknown labels or states fail',
                     'Only Table 6B used; Table 6A taxable-status subsets not stacked into it'],
    }
    (out / 'metadata.json').write_text(json.dumps(metadata, indent=2), encoding='utf-8')
    print(json.dumps(metadata, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--xlsx', type=Path, help='Downloaded official 2021-22 Table 6 workbook')
    parser.add_argument('--consumers', type=Path, help='Optional pipe-delimited tbl_consumer.csv or supplied Part 1 ZIP; postcode audit only')
    parser.add_argument('--compare', nargs='*', metavar='NAME=CSV',
                        help='Optional cleaned external tables with a postcode column (e.g. seifa=external_seifa/results/seifa_clean.csv) for joint coverage')
    parser.add_argument('--output', type=Path, help='Empty output directory; defaults to external_ato/results')
    run(parser.parse_args())
